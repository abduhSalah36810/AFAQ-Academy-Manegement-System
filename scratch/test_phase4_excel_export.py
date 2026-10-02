import asyncio
import io
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
TEST_DIRECTORY = tempfile.TemporaryDirectory(prefix="afaq-phase4-")
os.environ["TURSO_DATABASE_URL"] = ""
os.environ["TURSO_AUTH_TOKEN"] = ""
os.environ["SQLITE_DB_PATH"] = str(Path(TEST_DIRECTORY.name) / "phase4.db")

from database import create_tables, get_connection
from services.batch_attendance_service import BatchAttendanceService
from services.batch_enrollment_service import BatchEnrollmentService
from services.batch_service import BatchService
from services.bonus_service import BonusService
from services.course_service import CourseService
from services.grading_service import GradingService
from services.session_service import SessionService
from services import excel_service
import web_app
from starlette.requests import Request


class OfficialBatchWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        create_tables()
        cls.connection = get_connection()
        cursor = cls.connection.cursor()
        users = [
            (1, "مدير", "admin@afaq.admin.edu", "admin"),
            (2, "محاضر الدفعة", "instructor@afaq.instructor.edu", "instructor"),
            (10, "سارة علي", "sara@afaq.trainee.edu", "trainee"),
            (11, "محمود حسن", "mahmoud@afaq.trainee.edu", "trainee"),
            (12, "ليلى أحمد", "layla@afaq.trainee.edu", "trainee"),
        ]
        cursor.executemany(
            "INSERT INTO users (id, name, email, password_hash, role, active) VALUES (?, ?, ?, 'hash', ?, 1)",
            users,
        )
        cls.connection.commit()

        course_id = CourseService(cls.connection).create(
            "دورة تطوير البرمجيات", instructor_id=2
        )
        cls.batch_id = BatchService(cls.connection).create(
            course_id, "دفعة أكتوبر", instructor_id=2, start_date="2026-10-01"
        )
        sessions = SessionService(cls.connection)
        cls.session_one = sessions.create(
            cls.batch_id, "مقدمة", date="2026-10-02"
        )
        cls.session_two = sessions.create(
            cls.batch_id, "تطبيق عملي", date="2026-10-02"
        )
        cls.session_three = sessions.create(
            cls.batch_id, "مراجعة", date="2026-10-09"
        )

        enrollments = BatchEnrollmentService(cls.connection)
        for trainee_id in (10, 11, 12):
            enrollments.enroll(cls.batch_id, trainee_id, enrolled_by=1)

        attendance = BatchAttendanceService(cls.connection)
        attendance.record(cls.batch_id, 10, "2026-10-02", "present", session_id=cls.session_one)
        attendance.record(cls.batch_id, 10, "2026-10-09", "excused", session_id=cls.session_three)
        attendance.record(cls.batch_id, 11, "2026-10-02", "late")
        attendance.record(cls.batch_id, 12, "2026-10-09", "absent", session_id=cls.session_three)

        grading = GradingService(cls.connection)
        cls.component_exam = grading.create_component(cls.batch_id, "الاختبار", 40)
        cls.component_project = grading.create_component(cls.batch_id, "المشروع", 60)
        grading.set_score(cls.component_exam, 10, 80, recorded_by=1)
        grading.set_score(cls.component_project, 10, 90, recorded_by=1)
        grading.set_score(cls.component_exam, 11, 50, recorded_by=1)
        BonusService(cls.connection).award(10, 20, "تميز", awarded_by=1, batch_id=cls.batch_id)

    @classmethod
    def tearDownClass(cls):
        cls.connection.close()
        TEST_DIRECTORY.cleanup()

    def _loaded_workbook(self):
        batch = BatchService(self.connection).get_by_id(self.batch_id)
        sessions = SessionService(self.connection).get_by_batch(self.batch_id)
        attendance_service = BatchAttendanceService(self.connection)
        trainees = BatchEnrollmentService(self.connection).get_batch_trainees(self.batch_id)
        grading = GradingService(self.connection)
        components = grading.get_components(self.batch_id)
        score_rows = grading.get_all_scores_for_batch(self.batch_id)
        bonuses = BonusService(self.connection).get_batch_bonus_totals(self.batch_id)
        org_name = "أكاديمية آفاق"
        batch_info = {
            "academy_name": org_name,
            "course_name": batch[14],
            "batch_name": batch[2],
            "instructor_name": batch[4],
            "start_date": batch[6],
        }
        data = excel_service.build_official_batch_workbook(
            batch_info,
            sessions,
            trainees,
            attendance_service.get_batch_attendance_export_data(self.batch_id),
            components,
            score_rows,
            bonuses,
        )
        return load_workbook(io.BytesIO(data), data_only=True)

    def test_attendance_matrix_statuses_missing_values_and_same_day_safety(self):
        workbook = self._loaded_workbook()
        self.assertEqual(workbook.sheetnames, ["الحضور", "الدرجات"])
        sheet = workbook["الحضور"]
        self.assertEqual(sheet.max_column, 10)  # Two identity columns, 3 sessions, 5 summary columns.
        self.assertEqual(sheet["A1"].value, "أكاديمية آفاق")
        self.assertEqual(sheet["A4"].value, "اسم الأكاديمية")
        self.assertEqual(sheet["D6"].value, 3)  # Actual batch session count.
        self.assertEqual(sheet["B8"].value, "اسم المتدرب")
        self.assertEqual(sheet["C8"].value.splitlines()[0], "الجلسة 1")
        self.assertIn("مقدمة", sheet["C8"].value)
        rows_by_name = {
            sheet.cell(row, 2).value: row
            for row in range(9, 9 + 3)
        }
        sara_row = rows_by_name["سارة علي"]
        mahmoud_row = rows_by_name["محمود حسن"]
        layla_row = rows_by_name["ليلى أحمد"]
        self.assertEqual(sheet.cell(sara_row, 3).value, "حاضر")
        self.assertIsNone(sheet.cell(sara_row, 4).value)  # Same-day second session has no separate record.
        self.assertEqual(sheet.cell(sara_row, 5).value, "مستأذن")
        self.assertEqual(sheet.cell(mahmoud_row, 3).value, "غير محدد")
        self.assertEqual(sheet.cell(mahmoud_row, 4).value, "غير محدد")
        self.assertIsNone(sheet.cell(layla_row, 3).value)  # No attendance is left blank, not invented as absent.
        self.assertEqual(sheet.cell(sara_row, 6).value, 1)  # Present records.
        self.assertEqual(sheet.cell(sara_row, 9).value, 1)  # Excused records.
        self.assertEqual(sheet.cell(sara_row, 10).value, 50)  # 1 present / 2 attendance records.
        self.assertEqual(sheet.cell(sara_row, 10).number_format, '0.0"%"')
        self.assertEqual(sheet.cell(layla_row, 5).value, "غائب")
        notes = " ".join(
            str(sheet.cell(row, 1).value or "")
            for row in range(1, sheet.max_row + 1)
        )
        self.assertIn("قاعدة البيانات تحفظ سجل حضور واحدًا فقط", notes)
        self.assertIn("يتعذر إسناد سجل الحضور", notes)
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertEqual(sheet.freeze_panes, "C9")
        self.assertEqual(sheet.page_setup.orientation, "landscape")
        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)

    def test_grades_weights_final_bonus_and_capped_total(self):
        workbook = self._loaded_workbook()
        sheet = workbook["الدرجات"]
        self.assertEqual(sheet["D8"].value, "الاختبار\nالوزن: 40.0%")
        self.assertEqual(sheet["E8"].value, "المشروع\nالوزن: 60.0%")
        rows_by_email = {
            sheet.cell(row, 3).value: row
            for row in range(9, 12)
        }
        sara_row = rows_by_email["sara@afaq.trainee.edu"]
        mahmoud_row = rows_by_email["mahmoud@afaq.trainee.edu"]
        layla_row = rows_by_email["layla@afaq.trainee.edu"]
        self.assertEqual(sheet.cell(sara_row, 4).value, 80)
        self.assertEqual(sheet.cell(sara_row, 5).value, 90)
        self.assertEqual(sheet.cell(sara_row, 6).value, 86)
        self.assertEqual(sheet.cell(sara_row, 7).value, 20)
        self.assertEqual(sheet.cell(sara_row, 8).value, 100)
        self.assertEqual(sheet.cell(mahmoud_row, 6).value, 20)
        self.assertEqual(sheet.cell(mahmoud_row, 8).value, 20)
        self.assertEqual(sheet.cell(layla_row, 6).value, 0)
        self.assertEqual(sheet.cell(layla_row, 8).value, 0)

    def test_export_route_uses_one_batch_score_query_and_preserves_access(self):
        statements = []
        connection = get_connection()
        connection.set_trace_callback(statements.append)
        admin_request = Request({
            "type": "http", "method": "GET",
            "path": f"/admin/batches/{self.batch_id}/export-report",
            "headers": [],
            "session": {"user": {"id": 1, "role": "admin", "name": "مدير"}},
        })
        with patch.object(web_app, "get_db", return_value=connection):
            response = asyncio.run(
                web_app.export_official_batch_report(admin_request, self.batch_id)
            )
            payload = asyncio.run(self._read_stream(response))

        self.assertEqual(response.status_code, 200)
        self.assertIn("batch_" + str(self.batch_id) + "_report.xlsx", response.headers["content-disposition"])
        workbook = load_workbook(io.BytesIO(payload), data_only=True)
        self.assertEqual(workbook.sheetnames, ["الحضور", "الدرجات"])
        score_queries = [
            sql for sql in statements
            if "CROSS JOIN grading_components gc" in sql
            and "LEFT JOIN component_scores cs" in sql
        ]
        self.assertEqual(len(score_queries), 1)
        self.assertFalse(any("WHERE gc.batch_id = ?" in sql and "cs.trainee_id = ?" in sql for sql in statements))

    def test_legacy_export_urls_and_workbook_shapes_remain_available(self):
        route_paths = {route.path for route in web_app.app.routes}
        for prefix in ("/admin", "/instructor"):
            self.assertIn(f"{prefix}/batches/{{batch_id}}/export-attendance", route_paths)
            self.assertIn(f"{prefix}/batches/{{batch_id}}/export-grades", route_paths)

        old_attendance_rows = BatchAttendanceService(self.connection).get_batch_attendance(self.batch_id)
        old_attendance = [
            (row[1], row[2], row[3], row[5]) for row in old_attendance_rows
        ]
        attendance_book = load_workbook(io.BytesIO(
            excel_service.export_attendance("دفعة أكتوبر", old_attendance)
        ), data_only=True)
        self.assertEqual(attendance_book.sheetnames, ["Attendance"])

        components, _trainees, grade_rows = web_app._get_batch_grade_export_rows(
            self.connection, self.batch_id
        )
        grades_book = load_workbook(io.BytesIO(
            excel_service.export_grades("دفعة أكتوبر", components, grade_rows)
        ), data_only=True)
        self.assertEqual(grades_book.sheetnames, ["Grades"])

    def test_legacy_exports_support_ascii_and_arabic_batch_names(self):
        original = self.connection.execute(
            "SELECT name FROM batches WHERE id = ?", (self.batch_id,)
        ).fetchone()[0]
        routes = (
            (web_app.admin_export_attendance, "attendance", "Attendance"),
            (web_app.admin_export_grades, "grades", "Grades"),
        )
        try:
            for batch_name in ("October Batch", "دفعة أكتوبر"):
                self.connection.execute(
                    "UPDATE batches SET name = ? WHERE id = ?",
                    (batch_name, self.batch_id),
                )
                self.connection.commit()

                for route, prefix, expected_sheet in routes:
                    connection = get_connection()
                    request = Request({
                        "type": "http", "method": "GET",
                        "path": f"/admin/batches/{self.batch_id}/export-{prefix}",
                        "headers": [],
                        "session": {"user": {"id": 1, "role": "admin", "name": "مدير"}},
                    })
                    with patch.object(web_app, "get_db", return_value=connection):
                        response = asyncio.run(route(request, self.batch_id))
                        payload = asyncio.run(self._read_stream(response))

                    self.assertEqual(response.status_code, 200)
                    disposition = response.headers["content-disposition"]
                    disposition.encode("ascii")
                    self.assertIn('filename="', disposition)
                    self.assertIn("filename*=UTF-8''", disposition)
                    if batch_name == "دفعة أكتوبر":
                        self.assertIn(f"batch_{self.batch_id}.xlsx", disposition)
                        self.assertIn("%D8", disposition)
                    else:
                        self.assertIn("October_Batch.xlsx", disposition)
                    workbook = load_workbook(io.BytesIO(payload), data_only=True)
                    self.assertEqual(workbook.sheetnames, [expected_sheet])
        finally:
            self.connection.execute(
                "UPDATE batches SET name = ? WHERE id = ?",
                (original, self.batch_id),
            )
            self.connection.commit()

    def test_batch_grade_pages_reuse_one_score_query_and_keep_values(self):
        cases = (
            (web_app.admin_batch_grades, "/admin/batches/grades", {"id": 1, "role": "admin", "name": "مدير"}),
            (web_app.instructor_batch_grades, "/instructor/batches/grades", {"id": 2, "role": "instructor", "name": "محاضر"}),
        )
        for route, path, user in cases:
            statements = []
            connection = get_connection()
            connection.set_trace_callback(statements.append)
            request = Request({
                "type": "http", "method": "GET", "path": path,
                "headers": [], "session": {"user": user},
            })
            with patch.object(web_app, "get_db", return_value=connection), patch.object(
                web_app, "render", side_effect=lambda _request, _template, **context: context
            ):
                context = asyncio.run(route(request, self.batch_id))

            score_queries = [
                sql for sql in statements
                if "CROSS JOIN grading_components gc" in sql
                and "LEFT JOIN component_scores cs" in sql
            ]
            self.assertEqual(len(score_queries), 1)
            self.assertFalse(any("cs.trainee_id = ?" in sql for sql in statements))
            grades_by_id = {row[0]: row for row in context["grades_summary"]}
            self.assertEqual(grades_by_id[10][4:], (86.0, 20))
            self.assertEqual(grades_by_id[11][4:], (20.0, 0))
            self.assertEqual(grades_by_id[12][4:], (0.0, 0))

    def test_batch_grades_service_summary_uses_batch_score_data(self):
        statements = []
        connection = get_connection()
        connection.set_trace_callback(statements.append)
        summary = GradingService(connection).get_batch_grades_summary(self.batch_id)
        connection.close()

        score_queries = [
            sql for sql in statements
            if "CROSS JOIN grading_components gc" in sql
            and "LEFT JOIN component_scores cs" in sql
        ]
        self.assertEqual(len(score_queries), 1)
        self.assertFalse(any("cs.trainee_id = ?" in sql for sql in statements))
        self.assertEqual(
            {row[0]: row[3] for row in summary},
            {10: 86.0, 11: 20.0, 12: 0.0},
        )

    async def _read_stream(self, response):
        return b"".join([chunk async for chunk in response.body_iterator])


if __name__ == "__main__":
    unittest.main()
